#!/usr/bin/env python3
"""Build SpreadsheetBench tasks for the my-claude-harness evaluation framework.

SpreadsheetBench: real-world Excel manipulation.
Each task: agent edits an input .xlsx per an instruction, writes a result .xlsx.
Judge compares a specific cell range (answer_position) in a specific sheet
(answer_sheet) between the agent's output and the golden answer.

Lessons applied from OfficeQA cleaning:
  * README.md must carry the FULL instruction + explicit output contract (the
    agent only ever sees public/README.md as its task statement).
  * Paths in README are agent-relative (input at envs/data/, output to outputs/).
  * Judge is robust (argparse accepts all harness args; graceful errors).
  * Golden answer lives in evaluation/data/ (private, never shipped to agent).
  * Shared venv (with openpyxl) via symlink; backup existing task dir before overwrite.
"""
import argparse
import datetime
import json
import os
import pathlib
import shutil
import sys

REPO = pathlib.Path(__file__).resolve().parents[2]
BENCH = REPO / "bench_data" / "spreadsheetbench"
RAW_ROOT = BENCH / "raw" / "extracted" / "spreadsheetbench_verified_400"
DATASET_JSON = RAW_ROOT / "dataset.json"
SPREADSHEET_DIR = RAW_ROOT / "spreadsheet"
TASKS_DIR = REPO / "tasks" / "spreadsheetbench"
TASKS_DIR.mkdir(parents=True, exist_ok=True)

# Shared venv (created on server with openpyxl + pandas). On the local box this
# symlink target may not exist; that's fine — the harness resolves it server-side.
VENV_TARGET = os.environ.get(
    "SSB_VENV", str(TASKS_DIR / ".shared_venv")
)

OUTPUT_SCHEMA = {
    "type": "object",
    "required": ["output_file"],
    "properties": {
        "output_file": {
            "type": "string",
            "description": "Relative path (under outputs/) to the completed .xlsx file",
        }
    },
}

JUDGE_CODE = r'''#!/usr/bin/env python3
"""SpreadsheetBench judge.

Compares the agent's output .xlsx against the golden .xlsx over the target
cell range (answer_position) in the target sheet (answer_sheet).

Reward:
  1.0  -> every target cell matches the golden value exactly (after normalization)
  partial -> fraction of matching cells (only if >0 and <1), reported for feedback
  0.0  -> no output / unreadable / no cells match
"""
import argparse
import json
import os

def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--submission", required=True)
    p.add_argument("--cases")
    p.add_argument("--schema")
    p.add_argument("--metrics")
    p.add_argument("--eval-data")
    p.add_argument("--result", required=True)
    p.add_argument("--feedback-level", default="metric_status")
    return p.parse_args()


def write_result(path, obj):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w") as f:
        json.dump(obj, f)
    print(json.dumps(obj))


def norm_cell(v):
    """Normalize a cell value for robust comparison."""
    if v is None:
        return ""
    if isinstance(v, float):
        # Treat integral floats as ints; round others to avoid FP noise.
        if v == int(v):
            return str(int(v))
        return str(round(v, 6))
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def col_to_idx(col):
    idx = 0
    for ch in col:
        idx = idx * 26 + (ord(ch.upper()) - ord("A") + 1)
    return idx


def split_ref(cell):
    """'A3' -> ('A', 3)"""
    i = 0
    while i < len(cell) and cell[i].isalpha():
        i += 1
    return cell[:i], int(cell[i:])


def iter_range_cells(rng):
    """'A3:D32' -> list of (row, col_idx) 1-based tuples."""
    if ":" in rng:
        start, end = rng.split(":")
    else:
        start, end = rng, rng
    sc, sr = split_ref(start)
    ec, er = split_ref(end)
    c0, c1 = col_to_idx(sc), col_to_idx(ec)
    r0, r1 = sr, er
    if c0 > c1:
        c0, c1 = c1, c0
    if r0 > r1:
        r0, r1 = r1, r0
    cells = []
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cells.append((r, c))
    return cells


def main():
    args = parse_args()
    eval_data = args.eval_data
    meta_path = os.path.join(eval_data, "meta.json")
    golden_path = os.path.join(eval_data, "golden.xlsx")

    if not os.path.exists(meta_path) or not os.path.exists(golden_path):
        write_result(args.result, {
            "status": "error", "reward": 0.0,
            "feedback": "judge misconfigured: missing meta.json or golden.xlsx",
        })
        return

    with open(meta_path) as f:
        meta = json.load(f)
    answer_sheet = meta["answer_sheet"]
    answer_position = meta["answer_position"]

    try:
        import openpyxl
    except Exception as e:
        write_result(args.result, {
            "status": "error", "reward": 0.0,
            "feedback": "judge env missing openpyxl: %s" % e,
        })
        return

    # Locate the agent's output xlsx in the submission dir.
    sub = args.submission
    out_xlsx = None
    # Preferred: a results.json pointing to the file, else first .xlsx found.
    rj = os.path.join(sub, "results.json")
    if os.path.exists(rj):
        try:
            with open(rj) as f:
                data = json.load(f)
            cand = data.get("output_file")
            if cand:
                cand_path = os.path.join(sub, cand)
                if os.path.exists(cand_path):
                    out_xlsx = cand_path
        except Exception:
            pass
    if out_xlsx is None:
        for root, _dirs, files in os.walk(sub):
            for fn in sorted(files):
                if fn.lower().endswith(".xlsx"):
                    out_xlsx = os.path.join(root, fn)
                    break
            if out_xlsx:
                break

    if out_xlsx is None:
        write_result(args.result, {
            "status": "fail", "reward": 0.0,
            "feedback": "No .xlsx output found in submission. Write your completed "
                        "spreadsheet to outputs/ (e.g. outputs/result.xlsx).",
        })
        return

    try:
        gwb = openpyxl.load_workbook(golden_path, data_only=True)
        owb = openpyxl.load_workbook(out_xlsx, data_only=True)
    except Exception as e:
        write_result(args.result, {
            "status": "fail", "reward": 0.0,
            "feedback": "Could not read a workbook: %s" % e,
        })
        return

    if answer_sheet not in gwb.sheetnames:
        # Fall back to active sheet if the named sheet is absent in golden.
        gsheet = gwb.active
    else:
        gsheet = gwb[answer_sheet]
    if answer_sheet in owb.sheetnames:
        osheet = owb[answer_sheet]
    else:
        osheet = owb.active

    cells = iter_range_cells(answer_position)
    total = len(cells)
    match = 0
    diffs = []
    for (r, c) in cells:
        gv = norm_cell(gsheet.cell(row=r, column=c).value)
        ov = norm_cell(osheet.cell(row=r, column=c).value)
        if gv == ov:
            match += 1
        elif len(diffs) < 5:
            diffs.append("cell(r%d,c%d): got %r want %r" % (r, c, ov, gv))

    frac = match / total if total else 0.0
    if frac >= 1.0:
        write_result(args.result, {
            "status": "pass", "reward": 1.0,
            "feedback": "All %d target cells match." % total,
        })
    elif frac > 0:
        write_result(args.result, {
            "status": "fail", "reward": round(frac, 4),
            "feedback": "Matched %d/%d target cells. Sample diffs: %s" % (
                match, total, "; ".join(diffs)),
        })
    else:
        write_result(args.result, {
            "status": "fail", "reward": 0.0,
            "feedback": "No target cells match (0/%d). Sample diffs: %s" % (
                total, "; ".join(diffs)),
        })


if __name__ == "__main__":
    main()
'''


def build_one_task(row):
    tid_raw = str(row["id"])
    tid = "spreadsheetbench-" + tid_raw.replace("_", "-")
    task_dir = TASKS_DIR / tid

    src_folder = SPREADSHEET_DIR / pathlib.Path(row["spreadsheet_path"]).name
    if not src_folder.exists():
        print(f"    Warning: spreadsheet folder not found: {src_folder}")
        return None

    # Find the input (init) and golden xlsx (verified_400 => case number 1).
    init_files = sorted(src_folder.glob("*_init.xlsx"))
    golden_files = sorted(src_folder.glob("*_golden.xlsx"))
    if not init_files or not golden_files:
        print(f"    Warning: missing init/golden xlsx in {src_folder}")
        return None
    init_xlsx = init_files[0]
    golden_xlsx = golden_files[0]

    if task_dir.exists():
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = TASKS_DIR / f"{tid}.bak_{ts}"
        shutil.copytree(str(task_dir), str(backup))
        print(f"    Backup -> {backup.name}")
        shutil.rmtree(str(task_dir))

    (task_dir / "envs" / "data").mkdir(parents=True, exist_ok=True)
    (task_dir / "evaluation" / "data").mkdir(parents=True, exist_ok=True)
    (task_dir / "envs" / "runtime").mkdir(parents=True, exist_ok=True)

    # Input file the agent may edit -> envs/data/input.xlsx
    shutil.copy2(str(init_xlsx), str(task_dir / "envs" / "data" / "input.xlsx"))

    # Golden + metadata -> evaluation/data/ (private)
    shutil.copy2(str(golden_xlsx), str(task_dir / "evaluation" / "data" / "golden.xlsx"))
    # NOTE: ~2/3 of rows have NO 'answer_sheet' (single-sheet workbooks). We store
    # an empty string and the judge falls back to the active sheet in that case.
    answer_sheet = row.get("answer_sheet", "") or ""
    answer_position = row["answer_position"]
    meta = {
        "id": tid_raw,
        "answer_sheet": answer_sheet,
        "answer_position": answer_position,
        "data_position": row.get("data_position", ""),
        "instruction_type": row.get("instruction_type", ""),
    }
    with open(task_dir / "evaluation" / "data" / "meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)

    # task_manifest.json
    manifest = {
        "version": 1,
        "task_id": tid,
        "public_bundle": ["envs/data/"],
        "entrypoints": {
            "judge": "evaluation/judge.py",
            "output_schema": "evaluation/output_schema.json",
            "environment": "envs/env_manifest.json",
        },
    }
    with open(task_dir / "task_manifest.json", "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False)

    # env_manifest.json
    env_m = {
        "default_env": "runtime",
        "envs": {
            "runtime": {
                "python": {
                    "posix": "envs/runtime/.venv/bin/python",
                    "windows": "envs/runtime/.venv/Scripts/python.exe",
                }
            }
        },
    }
    with open(task_dir / "envs" / "env_manifest.json", "w", encoding="utf-8") as f:
        json.dump(env_m, f, indent=2, ensure_ascii=False)

    # README.md — the ONLY task statement the agent sees.
    instruction = row["instruction"]
    if answer_sheet:
        graded_line = (
            f"Your result is graded by comparing the cell range "
            f"**{answer_position}** on the sheet named **\"{answer_sheet}\"** "
            f"against the reference answer. Make sure that sheet exists in your "
            f"output and that those cells contain the correct values. Preserve "
            f"the sheet name exactly."
        )
    else:
        graded_line = (
            f"Your result is graded by comparing the cell range "
            f"**{answer_position}** (on the workbook's first/active sheet) "
            f"against the reference answer. Make sure those cells contain the "
            f"correct values."
        )
    readme = f"""# Task: Spreadsheet Manipulation

You are given an Excel workbook. Complete the manipulation described below, then
save the completed workbook as a new file.

## Instruction

{instruction}

## Input

The input workbook is available at `envs/data/input.xlsx`. Do not assume any
particular sheet is active — inspect the workbook first (sheet names, headers,
data ranges) before editing.

## What is graded

{graded_line}

## Output Format

1. Load `envs/data/input.xlsx`, perform the requested manipulation with a library
   such as `openpyxl` (available in the runtime environment).
2. Save the completed workbook to `outputs/result.xlsx`.
3. Also write `outputs/results.json` describing your output file:

```json
{{
  "output_file": "result.xlsx"
}}
```

The graded values are read from the cells listed above, so double-check them.

## IMPORTANT: how to submit

This task requires only a few steps. Work efficiently and do NOT keep exploring
the workbook indefinitely. As soon as you have written `outputs/result.xlsx` and
`outputs/results.json` with the correct values, you MUST call the
**`finalize_submission`** tool to submit your work for grading. Your answer is
only graded after you call `finalize_submission`; if you never call it, the run
fails. Provide a one-line summary and list `outputs/result.xlsx` in the files.
"""
    with open(task_dir / "README.md", "w", encoding="utf-8") as f:
        f.write(readme)

    # output_schema.json
    with open(task_dir / "evaluation" / "output_schema.json", "w", encoding="utf-8") as f:
        json.dump(OUTPUT_SCHEMA, f, indent=2, ensure_ascii=False)

    # rubric.txt
    rubric = (
        f"Task: SpreadsheetBench - {tid_raw}\n"
        f"Type: {row.get('instruction_type','')}\n"
        f"Answer sheet: {answer_sheet or '(active sheet)'}\n"
        f"Answer range: {answer_position}\n"
        f"Scoring: exact cell-value match over the answer range (normalized).\n"
    )
    with open(task_dir / "evaluation" / "rubric.txt", "w", encoding="utf-8") as f:
        f.write(rubric)

    # judge.py
    with open(task_dir / "evaluation" / "judge.py", "w", encoding="utf-8") as f:
        f.write(JUDGE_CODE)

    # Shared venv symlink
    venv_link = task_dir / "envs" / "runtime" / ".venv"
    if not venv_link.exists():
        try:
            os.symlink(VENV_TARGET, str(venv_link), target_is_directory=True)
        except Exception:
            pass

    return task_dir


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--max-tasks", type=int, default=None)
    p.add_argument("--start", type=int, default=0)
    p.add_argument("--only", type=str, default=None, help="build only this id (e.g. 13-1)")
    args = p.parse_args()

    if not DATASET_JSON.exists():
        print(f"X dataset.json not found: {DATASET_JSON}")
        print("  Did you extract the tarball? See adapters/spreadsheetbench/README.")
        sys.exit(1)

    with open(DATASET_JSON, encoding="utf-8") as f:
        rows = json.load(f)

    if args.only:
        rows = [r for r in rows if str(r["id"]) == args.only]
    else:
        rows = rows[args.start:]
        if args.max_tasks:
            rows = rows[: args.max_tasks]

    print(f"Building {len(rows)} SpreadsheetBench tasks...")
    built = 0
    for i, row in enumerate(rows):
        tid = "spreadsheetbench-" + str(row["id"]).replace("_", "-")
        print(f"  [{i+1}/{len(rows)}] {tid}")
        if build_one_task(row):
            built += 1
    print(f"\nDone. Built: {built}")


if __name__ == "__main__":
    main()
